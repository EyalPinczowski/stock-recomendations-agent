from pptx import Presentation

from portfolio_agent.models import RiskProfile
from portfolio_agent.pipeline import run_analyze
from portfolio_agent.providers.market_mock import MockMarketDataProvider
from portfolio_agent.providers.portfolio_file import FilePortfolioProvider
from portfolio_agent.report.presentation import build_presentation


def test_build_presentation_produces_valid_pptx(tmp_path):
    provider = FilePortfolioProvider("examples/portfolio.csv")
    market = MockMarketDataProvider()
    report = run_analyze(provider, market, RiskProfile(), state_dir=str(tmp_path))

    out = tmp_path / "report.pptx"
    path = build_presentation(report, "analyze", str(out))
    assert path.exists()

    prs = Presentation(str(path))
    slide_count = len(prs.slides._sldIdLst)
    # title + overall assessment + health + (rebalance?) + recs/holds (+warnings?)
    assert slide_count >= 4


def test_build_presentation_works_without_matplotlib(tmp_path, monkeypatch):
    """Charts are optional — on hosts where matplotlib isn't installed (Termux),
    the deck must still build, with allocation tables instead of charts."""
    import portfolio_agent.report.presentation as presentation_mod

    monkeypatch.setattr(presentation_mod, "CHARTS_AVAILABLE", False)
    monkeypatch.setattr(presentation_mod, "plt", None)

    provider = FilePortfolioProvider("examples/portfolio.csv")
    market = MockMarketDataProvider()
    report = run_analyze(provider, market, RiskProfile(), state_dir=str(tmp_path))

    out = tmp_path / "no_charts.pptx"
    path = build_presentation(report, "analyze", str(out))
    assert path.exists()
    assert len(Presentation(str(path)).slides._sldIdLst) >= 4

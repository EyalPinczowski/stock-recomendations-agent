"""Builds the .pptx deck — the primary output. Every slide pairs numbers/charts
with a plain-English explanation; built from whatever PortfolioReport/
ScreenReport/NewStockIdeasReport the pipeline produced, the same single source
of truth the other renderers use.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

# matplotlib is optional: it's the single hardest dependency to install on
# constrained hosts (Termux/Android in particular). Without it the deck still
# builds — the allocation charts are rendered as tables instead.
try:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    CHARTS_AVAILABLE = True
except ImportError:  # pragma: no cover - depends on host environment
    plt = None
    CHARTS_AVAILABLE = False

from portfolio_agent.models import (
    Action,
    NewStockIdeasReport,
    PortfolioReport,
    Recommendation,
    ScreenReport,
)

BG = RGBColor(0xFA, 0xFA, 0xF8)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
WARN = RGBColor(0xB0, 0x3A, 0x2B)
GOOD = RGBColor(0x2E, 0x7D, 0x32)
MAX_ACTION_SLIDES = 10


class DeckBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self._blank = self.prs.slide_layouts[6]

    def new_slide(self):
        slide = self.prs.slides.add_slide(self._blank)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG
        return slide

    def add_title(self, slide, text, sub=None):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = INK
        if sub:
            p2 = tf.add_paragraph()
            p2.text = sub
            p2.font.size = Pt(14)
            p2.font.color.rgb = MUTED

    def add_text(self, slide, left, top, width, height, text, size=14, color=INK, bold=False, italic=False):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.italic = italic
        return box

    def add_table(self, slide, left, top, width, height, headers, rows):
        n_rows, n_cols = len(rows) + 1, len(headers)
        gt = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
        table = gt.table
        for c, h in enumerate(headers):
            cell = table.cell(0, c)
            cell.text = h
            cell.text_frame.paragraphs[0].font.bold = True
            cell.text_frame.paragraphs[0].font.size = Pt(13)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT
            cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                cell.text = str(val)
                cell.text_frame.paragraphs[0].font.size = Pt(12)
                cell.text_frame.paragraphs[0].font.color.rgb = INK
        return table

    def add_chart_image(self, slide, fig, left, top, height):
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", transparent=True)
        plt.close(fig)
        buf.seek(0)
        slide.shapes.add_picture(buf, left, top, height=height)

    def add_callout_box(self, slide, left, top, width, height, text):
        shape = slide.shapes.add_shape(1, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xEF, 0xF3, 0xF8)
        shape.line.color.rgb = RGBColor(0xCC, 0xDA, 0xE8)
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(13)
        p.font.color.rgb = INK
        p.font.italic = True


def _action_color(action: Action) -> RGBColor:
    if action in (Action.BUY, Action.ADD):
        return GOOD
    if action in (Action.TRIM, Action.SELL):
        return WARN
    return MUTED


def _add_recommendation_slide(deck: DeckBuilder, rec: Recommendation):
    slide = deck.new_slide()
    deck.add_title(
        slide,
        f"{rec.ticker} — {rec.action.value.upper()}",
        sub=f"Conviction: {rec.conviction:.2f}"
        + (f"   ·   Reward:Risk {rec.reward_risk_ratio:.1f}" if rec.reward_risk_ratio else ""),
    )
    y = Inches(1.5)
    if rec.stop_take.stop_loss:
        deck.add_text(slide, Inches(0.6), y, Inches(5.6), Inches(0.4),
                       f"Stop-loss: {rec.stop_take.stop_loss:.2f} ({rec.stop_take.stop_basis})", size=14, color=WARN, bold=True)
        y = Inches(2.0)
    if rec.stop_take.take_profit:
        deck.add_text(slide, Inches(0.6), y, Inches(5.6), Inches(0.4),
                       f"Take-profit: {rec.stop_take.take_profit:.2f} ({rec.stop_take.target_basis})", size=14, color=GOOD, bold=True)

    deck.add_text(slide, Inches(0.6), Inches(2.7), Inches(5.6), Inches(0.35), "Rationale", size=14, bold=True, color=ACCENT)
    deck.add_text(slide, Inches(0.6), Inches(3.1), Inches(5.6), Inches(2.6), rec.rationale, size=13)

    if rec.review:
        deck.add_text(slide, Inches(6.6), Inches(2.7), Inches(6.0), Inches(0.35), "Reviewer's second opinion", size=14, bold=True, color=ACCENT)
        deck.add_callout_box(slide, Inches(6.6), Inches(3.1), Inches(6.0), Inches(2.6),
                              f"[{rec.review.outcome}] {rec.review.notes}" if rec.review.notes else f"[{rec.review.outcome}]")


def build_portfolio_deck(report: PortfolioReport) -> Presentation:
    deck = DeckBuilder()

    # Title
    slide = deck.new_slide()
    deck.add_text(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.0), "Portfolio Report", size=40, bold=True)
    deck.add_text(slide, Inches(0.8), Inches(3.3), Inches(11.7), Inches(0.5),
                   f"Total value: ${report.total_value:,.2f}", size=20, color=ACCENT, bold=True)
    deck.add_text(slide, Inches(0.8), Inches(3.85), Inches(11.7), Inches(0.4),
                   f"As of {report.generated_at:%Y-%m-%d %H:%M} UTC", size=13, color=MUTED)
    deck.add_text(slide, Inches(0.8), Inches(4.25), Inches(11.7), Inches(0.4),
                   f"Risk profile: {report.risk_profile.name} "
                   f"({report.risk_profile.target_conservative_pct:.0%}/{report.risk_profile.target_aggressive_pct:.0%} "
                   f"conservative/aggressive, min reward:risk {report.risk_profile.min_reward_risk_ratio})",
                   size=12, color=MUTED, italic=True)

    # Overall assessment
    slide = deck.new_slide()
    deck.add_title(slide, "Overall Assessment")
    deck.add_text(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(4.5), report.overall_assessment or "", size=17)

    # Portfolio health
    slide = deck.new_slide()
    deck.add_title(slide, "Portfolio Health")
    h = report.health
    y = Inches(1.3)
    if h.sharpe_ratio is not None:
        deck.add_text(slide, Inches(0.6), y, Inches(3.8), Inches(0.5), f"Sharpe ratio: {h.sharpe_ratio:.2f}", size=18, bold=True, color=ACCENT)
        y = Inches(1.85)
    if h.volatility_annualized is not None:
        deck.add_text(slide, Inches(0.6), y, Inches(3.8), Inches(0.4), f"Annualized volatility: {h.volatility_annualized:.0%}", size=14)
        y = Inches(2.3)
    if h.max_drawdown_pct is not None:
        deck.add_text(slide, Inches(0.6), y, Inches(3.8), Inches(0.4), f"Max drawdown (1y): {h.max_drawdown_pct:.0%}", size=14, color=WARN)

    if h.bucket_allocation:
        if CHARTS_AVAILABLE:
            fig, ax = plt.subplots(figsize=(2.9, 2.9))
            labels = list(h.bucket_allocation.keys())
            values = [v * 100 for v in h.bucket_allocation.values()]
            ax.pie(values, labels=[f"{name}\n{v:.0f}%" for name, v in zip(labels, values)], colors=["#2B6CB0", "#B0742B"][: len(labels)], textprops={"fontsize": 9})
            ax.set_title(
                f"Bucket allocation (target {report.risk_profile.target_conservative_pct:.0%}/{report.risk_profile.target_aggressive_pct:.0%})",
                fontsize=9,
            )
            deck.add_chart_image(slide, fig, Inches(4.6), Inches(1.3), Inches(2.9))
        else:
            targets = {
                "conservative": report.risk_profile.target_conservative_pct,
                "aggressive": report.risk_profile.target_aggressive_pct,
            }
            deck.add_table(
                slide, Inches(4.6), Inches(1.3), Inches(3.4), Inches(0.4 * (len(h.bucket_allocation) + 1)),
                headers=["Bucket", "Current", "Target"],
                rows=[
                    [bucket, f"{pct:.0%}", f"{targets.get(bucket, 0):.0%}"]
                    for bucket, pct in h.bucket_allocation.items()
                ],
            )

    if h.sector_allocation:
        if CHARTS_AVAILABLE:
            fig, ax = plt.subplots(figsize=(4.2, 2.9))
            sectors = list(h.sector_allocation.keys())
            pct = [v * 100 for v in h.sector_allocation.values()]
            colors = ["#B03A2B" if p > report.risk_profile.max_sector_pct * 100 else "#2B6CB0" for p in pct]
            ax.barh(sectors, pct, color=colors)
            ax.axvline(report.risk_profile.max_sector_pct * 100, color="#5A5A5A", linestyle="--", linewidth=1)
            ax.set_xlabel("% of portfolio", fontsize=9)
            ax.set_title(f"Sector concentration (cap {report.risk_profile.max_sector_pct:.0%})", fontsize=9)
            ax.tick_params(labelsize=8)
            deck.add_chart_image(slide, fig, Inches(7.6), Inches(1.3), Inches(2.9))
        else:
            cap = report.risk_profile.max_sector_pct
            deck.add_table(
                slide, Inches(8.3), Inches(1.3), Inches(4.4), Inches(0.4 * (len(h.sector_allocation) + 1)),
                headers=["Sector", "Weight", f"Over {cap:.0%} cap?"],
                rows=[
                    [sector, f"{pct:.0%}", "yes" if pct > cap else ""]
                    for sector, pct in sorted(h.sector_allocation.items(), key=lambda kv: kv[1], reverse=True)
                ],
            )

    if h.concentration_warnings:
        deck.add_text(slide, Inches(0.6), Inches(4.6), Inches(12.1), Inches(0.8),
                       "  ".join(f"⚠ {w}" for w in h.concentration_warnings), size=13, color=WARN, bold=True)

    # Rebalance suggestions
    if report.rebalance_suggestions:
        slide = deck.new_slide()
        deck.add_title(slide, "Rebalance Suggestions")
        deck.add_table(
            slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.5 * (len(report.rebalance_suggestions) + 1)),
            headers=["Kind", "Label", "Current", "Target", "Drift", "Why"],
            rows=[
                [s.kind, s.label, f"{s.current_pct:.1%}", f"{s.target_pct:.1%}", f"{s.drift_pct:+.1%}", s.action_summary]
                for s in report.rebalance_suggestions
            ],
        )

    # Recommendation slides for actionable items, capped
    actionable = [r for r in report.holding_recommendations if r.action != Action.HOLD]
    plain_holds = [r for r in report.holding_recommendations if r.action == Action.HOLD]
    for rec in actionable[:MAX_ACTION_SLIDES]:
        _add_recommendation_slide(deck, rec)

    # Summary table for plain holds (and any overflow actionable items)
    overflow = actionable[MAX_ACTION_SLIDES:]
    summary_rows = plain_holds + overflow
    if summary_rows:
        slide = deck.new_slide()
        deck.add_title(slide, "Other Holdings", sub="No slide-worthy signal change" if plain_holds else "Additional recommendations")
        deck.add_table(
            slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(0.4 * (len(summary_rows) + 1)),
            headers=["Ticker", "Action", "Conviction", "Note"],
            rows=[[r.ticker, r.action.value.upper(), f"{r.conviction:.2f}", r.rationale[:80]] for r in summary_rows],
        )

    # Warnings
    if report.warnings:
        slide = deck.new_slide()
        deck.add_title(slide, "Warnings")
        for i, w in enumerate(report.warnings):
            deck.add_text(slide, Inches(0.6), Inches(1.4 + i * 0.5), Inches(12.1), Inches(0.5), f"• {w}", size=14, color=WARN)

    return deck.prs


def build_screen_deck(report: ScreenReport) -> Presentation:
    deck = DeckBuilder()
    slide = deck.new_slide()
    deck.add_text(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.0), "Short-Term Screen", size=40, bold=True)
    deck.add_text(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.4),
                   f"As of {report.generated_at:%Y-%m-%d %H:%M} UTC  ·  {len(report.candidates)} candidates", size=14, color=MUTED)

    if report.candidates:
        slide = deck.new_slide()
        deck.add_title(slide, "Top Candidates")
        rows = sorted(report.candidates, key=lambda c: c.setup_score, reverse=True)
        deck.add_table(
            slide, Inches(0.6), Inches(1.4), Inches(12.1), Inches(0.4 * (len(rows) + 1)),
            headers=["Ticker", "Score", "Horizon", "Entry Zone", "Stop", "Target"],
            rows=[
                [c.ticker, f"{c.setup_score:.2f}", f"{c.horizon_days}d",
                 f"{c.entry_zone_low:.2f}-{c.entry_zone_high:.2f}",
                 f"{c.stop_take.stop_loss:.2f}" if c.stop_take.stop_loss else "-",
                 f"{c.stop_take.take_profit:.2f}" if c.stop_take.take_profit else "-"]
                for c in rows
            ],
        )

    for c in sorted(report.candidates, key=lambda c: c.setup_score, reverse=True)[:MAX_ACTION_SLIDES]:
        slide = deck.new_slide()
        deck.add_title(slide, f"{c.ticker}", sub=f"Setup score {c.setup_score:.2f}  ·  {c.horizon_days}-day horizon")
        deck.add_text(slide, Inches(0.6), Inches(1.5), Inches(11.7), Inches(0.4),
                       f"Entry: {c.entry_zone_low:.2f}-{c.entry_zone_high:.2f}   Stop: {c.stop_take.stop_loss}   Target: {c.stop_take.take_profit}",
                       size=14, bold=True)
        deck.add_text(slide, Inches(0.6), Inches(2.1), Inches(11.7), Inches(2.5), c.rationale, size=13)

    if report.warnings:
        slide = deck.new_slide()
        deck.add_title(slide, "Warnings")
        for i, w in enumerate(report.warnings):
            deck.add_text(slide, Inches(0.6), Inches(1.4 + i * 0.5), Inches(12.1), Inches(0.5), f"• {w}", size=14, color=WARN)

    return deck.prs


def build_newstocks_deck(report: NewStockIdeasReport) -> Presentation:
    deck = DeckBuilder()
    slide = deck.new_slide()
    deck.add_text(slide, Inches(0.8), Inches(2.6), Inches(11.7), Inches(1.0), "New Stock Ideas", size=40, bold=True)
    deck.add_text(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.4),
                   f"As of {report.generated_at:%Y-%m-%d %H:%M} UTC", size=14, color=MUTED)
    if report.gaps_identified:
        deck.add_text(slide, Inches(0.8), Inches(4.0), Inches(11.7), Inches(1.0),
                       "Scanning for: " + "; ".join(report.gaps_identified), size=13, color=MUTED, italic=True)

    for s in report.suggestions:
        slide = deck.new_slide()
        deck.add_title(slide, f"{s.ticker} — {s.bucket.value}", sub=f"Sector: {s.sector or 'Unknown'}  ·  Score {s.composite_score:.2f}  ·  Suggested {s.suggested_allocation_pct:.1%}")
        deck.add_text(slide, Inches(0.6), Inches(1.5), Inches(5.6), Inches(0.35), "Why it fits", size=14, bold=True, color=ACCENT)
        deck.add_text(slide, Inches(0.6), Inches(1.9), Inches(5.6), Inches(1.5), s.fit_reason, size=13)
        deck.add_text(slide, Inches(0.6), Inches(3.5), Inches(5.6), Inches(0.35), "Rationale", size=14, bold=True, color=ACCENT)
        deck.add_text(slide, Inches(0.6), Inches(3.9), Inches(5.6), Inches(2.0), s.rationale, size=13)
        if s.review:
            deck.add_text(slide, Inches(6.6), Inches(1.5), Inches(6.0), Inches(0.35), "Reviewer's second opinion", size=14, bold=True, color=ACCENT)
            deck.add_callout_box(slide, Inches(6.6), Inches(1.9), Inches(6.0), Inches(2.0),
                                  f"[{s.review.outcome}] {s.review.notes}" if s.review.notes else f"[{s.review.outcome}]")

    if report.warnings:
        slide = deck.new_slide()
        deck.add_title(slide, "Warnings")
        for i, w in enumerate(report.warnings):
            deck.add_text(slide, Inches(0.6), Inches(1.4 + i * 0.5), Inches(12.1), Inches(0.5), f"• {w}", size=14, color=WARN)

    return deck.prs


_BUILDERS = {
    "analyze": build_portfolio_deck,
    "screen": build_screen_deck,
    "newstocks": build_newstocks_deck,
}


def build_presentation(report, report_kind: str, output_path: str | None = None) -> Path:
    builder = _BUILDERS[report_kind]
    prs = builder(report)
    if output_path is None:
        output_path = f"{report_kind}_{datetime.now(UTC):%Y%m%d_%H%M%S}.pptx"
    path = Path(output_path)
    prs.save(path)
    return path
